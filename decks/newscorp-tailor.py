#!/usr/bin/env python3
"""Tailor the Sobrynth WELL deck for the News Corp HR leaders briefing."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

SRC = "stage1.pptx"
OUT = "Sobrynth-NewsCorp-HR-Briefing.pptx"

INK    = RGBColor(0x24, 0x1D, 0x3B)
PURPLE = RGBColor(0x59, 0x46, 0x8D)
ORANGE = RGBColor(0xE0, 0x8A, 0x2E)
AMBER  = RGBColor(0xFF, 0xAD, 0x5B)
RED    = RGBColor(0xC8, 0x5F, 0x5F)
RUST   = RGBColor(0x9E, 0x48, 0x3F)
TEAL   = RGBColor(0x00, 0x77, 0x88)
BODY   = RGBColor(0x42, 0x3F, 0x4B)
MUTED  = RGBColor(0x7C, 0x77, 0x88)


def find(slide, name, contains=None):
    """Find a shape by name; disambiguate duplicates with a text fragment."""
    hits = [s for s in slide.shapes if s.name == name]
    if contains is not None:
        hits = [s for s in hits
                if s.has_text_frame and contains in s.text_frame.text]
    if not hits:
        raise KeyError(f"no shape {name!r} (contains={contains!r})")
    return hits[0]


def set_text(shape, text):
    """Replace all text, keeping the first run's character formatting."""
    tf = shape.text_frame
    paras = tf.paragraphs
    p0 = paras[0]
    if not p0.runs:
        raise ValueError(f"{shape.name!r} has no runs to inherit formatting from")
    p0.runs[0].text = text
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)


def set_lines(shape, lines):
    """Replace text with several paragraphs, cloning paragraph 0's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        raise ValueError(f"{shape.name!r} has no runs")
    template = copy.deepcopy(p0._p)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    set_text(shape, lines[0])
    parent = p0._p.getparent()
    prev = p0._p
    for line in lines[1:]:
        new = copy.deepcopy(template)
        prev.addnext(new)
        prev = new
        from pptx.text.text import _Paragraph
        para = _Paragraph(new, tf)
        para.runs[0].text = line
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
    return shape


def box(shape, left=None, top=None, width=None, height=None):
    if left is not None:   shape.left = Inches(left)
    if top is not None:    shape.top = Inches(top)
    if width is not None:  shape.width = Inches(width)
    if height is not None: shape.height = Inches(height)


def style(shape, size=None, color=None, bold=None):
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if size is not None:  r.font.size = Pt(size)
            if bold is not None:  r.font.bold = bold
            if color is not None: r.font.color.rgb = color


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def drop(shape):
    shape._element.getparent().remove(shape._element)


prs = Presentation(SRC)
S = prs.slides

# ----------------------------------------------------------------- slide 1
s = S[0]
set_text(find(s, "Text 2"),
         "24/7 peer recovery navigation for employees and their families — "
         "activated company by company, alongside the EAP and behavioral "
         "health benefits you already offer.")
box(find(s, "Text 2"), width=7.55, height=1.05)  # clear the maze watermark

tb = s.shapes.add_textbox(Inches(0.67), Inches(5.42), Inches(9.20), Inches(0.35))
tb.name = "Audience Line"
tb.text_frame.word_wrap = True
r = tb.text_frame.paragraphs[0].add_run()
r.text = ("PREPARED FOR NEWS CORP  ·  HR LEADERSHIP BRIEFING  "
          "·  FRIDAY, SEPTEMBER 18, 2026")
r.font.name, r.font.size, r.font.bold = "Roboto", Pt(12), True
r.font.color.rgb = AMBER

notes(s, """Jess

I want to start with the word on this slide that matters most, and it isn't 'finally.' It's alongside.

I know what this room is managing. Across News Corp you have an EAP. You have behavioral health coverage. Some of your companies have a digital mental health app, a musculoskeletal vendor, a fertility benefit, a navigation layer on top of all of it — and somebody above you asking why the stack keeps growing.

So I'm not here to add a fourteenth vendor. I'm here about one condition that every single thing in your stack technically covers and functionally does not reach.

I'm Jess Ferretti, VP of Sales at Sobrynth, and before anything else you should know I'm in recovery myself. I say that in rooms like this because the reason this condition doesn't show up in your utilization reports is that people like me learned very early not to raise our hand at work. That's not a benefits design problem. It's a disclosure problem.

One more framing note for this room specifically. You are not one workforce. You are a newsroom, a book publisher, a digital real estate business, a print and distribution operation, and a corporate center — across several countries. One EAP gets procured centrally and lands completely differently in each of them. Hold that thought, because it's the whole reason I'm asking for one company and ninety days, not a group-wide rollout.""")

# ----------------------------------------------------------------- slide 2
s = S[1]
set_text(find(s, "Text 5"),
         "Often, you can't. Many people with SUD are employed and "
         "high-performing — filing on deadline, hitting their numbers — and it "
         "stays invisible until an absence, a safety incident, a workers' comp "
         "claim, or a crisis.")
set_text(find(s, "Text 13"),
         "An EAP may technically cover SUD, but that doesn't mean employees will "
         "use it for SUD — or use it early. And when one EAP is procured "
         "centrally for very different workforces, the gap widens. Substance "
         "use needs an engagement strategy built for each culture, not one "
         "group-wide contract.")

notes(s, """Jess

There are three things I hear constantly, and I'd love a show of hands — or put it in the chat — for who has heard these before.

One — 'you can tell when someone has a substance use problem.' If you met me with my two kids, my SUV and my very quintessential house in the suburbs, would you know I spent the vast majority of my life struggling with drug use? Probably not. The picture people carry is somebody visibly falling apart. The actual population is functioning, employed, filing on deadline, hitting most of their numbers, and completely invisible until there's an absence, an incident, a claim, or a crisis. By the time it's visible on your dashboard, you're years in.

Two — 'if they really wanted help, they'd get it.' This is the one that sounds like common sense and isn't. Fear of stigma, fear of job consequences, cost, confidentiality — and the big one nobody designs for: they're not ready for formal treatment yet. Support has to exist before someone is ready, or it doesn't get used.

Three — and this is the one I'd flag hardest for this group — 'our EAP already handles substance use.' Here's the News Corp version of that problem. One EAP, procured centrally, lands in a newsroom, a book publisher, a digital real estate sales floor and a print plant. Those are four completely different disclosure cultures. A single group-wide engagement strategy will underperform in at least three of them, and you will never see it in the reporting, because the reporting is blended.""")

# ----------------------------------------------------- slide 3 (NEW: media)
s = S[2]
set_text(find(s, "Text 0"), "THE MEDIA WORKFORCE")
set_text(find(s, "Text 1"), "What the Data Says About People Who Work in Media")

CARDS = [
    # card bg, eyebrow, stat, caption, accent
    ("Shape 2",  "Text 3",  "Text 4",  "Text 5",  PURPLE,
     "SAMHSA · BY INDUSTRY", "3rd",
     "Arts, entertainment and recreation ranks 3rd of 19 industries for "
     "substance use disorder — 12.9% against 9.5% across all full-time workers"),
    ("Shape 6",  "Text 7",  "Text 8",  "Text 9",  ORANGE,
     "SAMHSA · BY INDUSTRY", "13.7%",
     "Past-month illicit drug use in arts and entertainment — second highest "
     "of any industry, against an 8.6% all-industry average"),
    ("Shape 10", "Text 11", "Text 12", "Text 13", RED,
     "TRAUMA EXPOSURE", "86–100%",
     "Of journalists report witnessing a traumatic event on the job. PTSD runs "
     "4–13%, and 28.6% lifetime among war correspondents"),
    ("Shape 14", "Text 15", "Text 16", "Text 17", RUST,
     "BURNOUT", "8 in 10",
     "US journalists report burnout or chronic stress in the past year. 38% say "
     "their mental health declined; half considered leaving"),
    ("Shape 18", "Text 19", "Text 20", "Text 21", TEAL,
     "IRREGULAR HOURS", "25%",
     "More alcohol consumed by night-shift workers than day-schedule workers. "
     "Rotating and night work is linked to binge drinking"),
    ("Shape 22", "Text 23", "Text 24", "Text 25", PURPLE,
     "JOB INSECURITY", "17,163",
     "Media job cuts announced in 2025, up 15% year over year. Sustained job "
     "insecurity is a documented driver of substance use"),
]

ROW_TOP = {0: 1.78, 1: 4.24}
CARD_H = 2.30
for i, (bg, eb, stat, cap, accent, eb_t, stat_t, cap_t) in enumerate(CARDS):
    top = ROW_TOP[i // 3]
    box(find(s, bg), top=top, height=CARD_H)

    e = find(s, eb)
    set_text(e, eb_t)
    style(e, size=12, color=accent, bold=True)
    box(e, top=top + 0.24, height=0.30)

    st = find(s, stat)
    set_text(st, stat_t)
    style(st, size=32, color=accent, bold=True)
    box(st, top=top + 0.60, height=0.55)

    c = find(s, cap)
    set_text(c, cap_t)
    style(c, size=12, color=BODY)
    box(c, top=top + 1.22, height=0.90)

foot = find(s, "Text 26")
set_text(foot,
         "Sources: SAMHSA NSDUH, Substance Use and Substance Use Disorder by "
         "Industry · Feinstein et al. 2002 · Dart Center for Journalism and "
         "Trauma · Muck Rack State of Journalism 2025 · Media Resilience "
         "Network 2026 · European Addiction Research, shiftwork systematic "
         "review · Challenger, Gray & Christmas 2025")
style(foot, size=10, color=MUTED)
box(foot, top=6.68, height=0.55, width=12.36)

notes(s, """Jess

This is the slide I added for this room, because the first question I always get is 'sure, but is that our people?'

Two things about the industry data. The top row is measured prevalence. SAMHSA breaks substance use disorder out by industry, and arts, entertainment and recreation comes in third of nineteen — 12.9% against a 9.5% average. Illicit drug use is second highest of any industry in the country. That's the entertainment-adjacent side of media, and it is a real signal, but I want to be straight with you: it is not a clean proxy for a newsroom or a book publisher. Take it as directional.

The bottom row is why I think the exposure is broader than the industry code suggests. It's three mechanisms, and every one of them is present somewhere in this company.

Trauma exposure. Between 86 and 100% of journalists report witnessing a traumatic event on the job. Not covering one. Witnessing one. PTSD runs four to thirteen percent, and among war correspondents it's 28.6% lifetime — four to five times the general population. Feinstein also found war journalists self-report drinking more than other journalists. If you have foreign desks, crime desks, or anyone who handles graphic user-generated content for a living, that's your population.

Burnout. Eight in ten US journalists report burnout or chronic stress in the past year. Thirty-eight percent say their mental health declined. Half have considered leaving.

Irregular hours. Night-shift workers consume about 25% more alcohol than day-schedule workers, and rotating shifts are independently associated with binge drinking. News does not keep business hours. Neither do print plants or distribution.

And then job insecurity underneath all of it — 17,163 media job cuts announced in 2025, up fifteen percent.

I'm not telling you News Corp has a drinking problem. I'm telling you that you employ, at scale, almost every occupational risk factor the literature identifies for this condition. And the one benefit you have pointed at it is the one your people are least likely to call.""")

# ------------------------------------------------- slide 5 (SUD at scale)
s = S[4]
hl = find(s, "Text 31")
set_text(hl, "At ~23,000 employees, 1 in 10 is roughly 2,300 people.")
notes(s, """Jess

Run this against your own census while we talk through it.

70% of adults with a substance use disorder hold active jobs. They're on your headcount. They enrolled during open enrollment.

1 in 10 of your employees has untreated SUD right now. News Corp runs roughly 23,000 people worldwide, so that's on the order of 2,300. And here's the number I want you to hold: 4 in 5 of them are never treated. That's about 1,800 people with nothing in place.

Now — 4 in 5 never treated. Not 4 in 5 uncovered. Every one of those people has coverage. They have your medical plan. They have your EAP. Coverage was never the problem. Reach is.

Over $8,800 per affected employee per year, up 30% in three years. I'm deliberately not going to multiply that across your whole census and put a headline number on the screen, because your operating companies are on different plans in different countries and the blended figure would be fiction. Run it per company. That's ask number one at the end.

And it isn't a projection — it's already in your numbers. It's spread across absence, turnover, disability, medical and pharmacy, and never once coded to substance use.""")

# --------------------------------------------- slide 6 (chain, re-cut)
s = S[5]
set_text(find(s, "Text 2"),  "A missed deadline")
set_text(find(s, "Text 8"),  "A relationship, quietly strained")
set_text(find(s, "Text 14"), "A resignation no one probed")
set_text(find(s, "Text 20"), "An incident in the plant or the field")
set_text(find(s, "Text 26"), "A body that's had enough")
set_text(find(s, "Text 18"),
         "6–9 months of salary: recruiting, onboarding, ramp-up to full "
         "productivity (SHRM)")

notes(s, """Jess

This is the slide to photograph.

This is not five employees. It's one person, moving left to right, over about three to five years. Every stage has a price tag, and every price tag lands somewhere in your world.

Stage one, a missed deadline. $15,000 in lost output. Nobody does anything — because a missed deadline isn't a crisis, it's a bad quarter. That's the whole problem: stage one never generates a ticket.

Stage two, a relationship quietly strains. $20,000 per leave event — continued premium plus backfill. That one is already crossing your desk as a leave request, and nobody codes it to substance use.

Stage three — and I re-cut this one for your workforce. In a plant it's a failed screening. In a newsroom or a publishing house it's a resignation nobody probed. Someone good leaves, the exit interview says 'pursuing other opportunities,' and it costs you $30,000 to $45,000 to replace them. Notice this is usually where employers finally act, and acting means losing them. You didn't resolve the exposure. You transferred it to another employer and paid replacement cost for the privilege.

Stage four, an incident. $47,316, average lost-time workers' comp claim. And I want to be specific with this room, because most people hear 'media' and picture a desk. You run print plants. You run distribution fleets. You put camera and field crews in unpredictable places. That is a safety-sensitive population inside a company that doesn't usually think of itself as a safety-sensitive employer.

And stage five, a body that's had enough — $485,320, the average cost of a single cancer diagnosis to a self-funded plan.

Someone just had to catch stage one.""")

# ------------------------------------ slide 10 (solution: per-company code)
s = S[9]
nav = find(s, "Text 14")
set_lines(nav, [
    "Dedicated line and unique code for each operating company",
    "Concierge support",
    "Free + community resources through to clinical care",
    "Employees + families",
])

# -------------------------------------------- slide 11 (90-day activation)
s = S[10]
set_text(find(s, "Text 2"),
         "A guided path from kickoff to launch, run company by company — so each "
         "masthead and business gets an activation built for its own culture, "
         "not a group-wide template.")

# ------------------------------------------------- slide 14 (Gardner case)
s = S[13]
set_text(find(s, "Text 2"),
         "250 employees. Minnesota commercial construction. The hardest room "
         "we have ever walked into.")
set_text(find(s, "Text 4"), "EMPLOYEE AWARENESS AT 45 DAYS")  # typo in source

# ------------------------------------------------- slide 15 (Olympia case)
s = S[14]
set_text(find(s, "Text 2"),
         "2,000+ employees. 42 locations. The closest analogue to a "
         "multi-company, multi-site footprint.")
notes(s, """Jess

Olympia is the case study I'd anchor on for News Corp, because the shape of the workforce is the closest analogue we have. Two thousand employees, forty-two locations, distributed, multiple operating cultures, one program.

Treatment-seeking tripled. And this is the number I want you to sit with: their health spend on this did not go up. Flat dollars year over year, while three times as many people got help. That's an independent three-year retrospective claims analysis run by a top-five health plan, not our own number.

Ninety-one percent workforce awareness within ninety days. Fifty-seven-plus peer coach contacts. And that's the floor, not the ceiling, because the claims data only covers employees actually on the health plan. The rest of the workforce has the same access and isn't counted here.

The reason I show you Gardner and Olympia together: Gardner proves it works where asking is hardest, Olympia proves it scales across sites without the spend following it.""")

# ----------------------------------------------- slide 16 (ROI, re-scaled)
s = S[15]
set_text(find(s, "Subtitle"),
         "Modeled at News Corp scale using the same per-employee economics. "
         "The multiplier will vary by workforce, the pattern will not.")
set_text(find(s, "BannerSubCaption"),
         "2,000-employee operating company · $100,000 investment")
set_text(find(s, "Card1Number"), "$100K")
set_text(find(s, "Card1Caption"),
         "Annual Sobrynth investment for one 2,000-person operating company")
set_text(find(s, "Card3Number"), "$864K")
set_text(find(s, "Card4Number"), "$204K")
set_text(find(s, "Closing"),
         "$1.07M in impact at one operating company. ~$12.3M across the group.")
set_text(find(s, "Support"),
         "Scaled linearly from Sobrynth's per-employee model using its own "
         "utilization data and independent peer-support research. The "
         "group-wide figure is illustrative, not a quote.")

notes(s, """Jess

Same model Sobrynth uses everywhere, re-scaled to your world so nobody has to do arithmetic in their head.

One two-thousand-person operating company invests a hundred thousand dollars a year. Engagement lands at six percent, three times the typical EAP benchmark, matched to what we actually see across our customer base. That avoids treatment cost and retains talent worth $864,000, and recovers $204,000 in productivity from employees who get, and stay, sober. Call it $1.07 million in impact against a hundred thousand dollar investment.

Carry the same per-employee economics across roughly 23,000 employees and you get about $12.3 million.

Two caveats I'd rather say than have someone catch. First, that group number is a straight linear scale — it's illustrative, not a quote, and your operating companies sit on different plans in different countries. Second, the savings side is modeled, not billed. The number I'd actually defend to a CFO is Olympia's: three times the treatment-seeking, flat dollars, verified by a health plan's own claims analysis.

So treat the $12.3 million as the shape of the prize, and the one-company pilot as how you find out what it's really worth here.""")

# ------------------------------------------------ slide 17 (next steps)
s = S[16]
set_text(find(s, "Text 5"),
         "before your next benefits review — one number per operating company, "
         "not one blended number for the group.")
box(find(s, "Text 5"), width=9.00)
set_text(find(s, "Text 8"),
         "Ask your EAP vendor for SUD-specific utilization, broken out by "
         "operating company.")
set_text(find(s, "Text 9"),
         "Blended, group-level reporting hides the variation. That variation "
         "is the entire business case.")
box(find(s, "Text 9"), width=9.20)

notes(s, """Jess

Three things, and none of them cost anything or require a decision.

One — run your workforce through the Substance Use Cost Calculator before your next benefits review. Ten minutes. But do it per operating company, not once for News Corp. A blended number across a newsroom, a publisher and a print operation is a number you can't act on.

Two — and this is the one I'd actually do first. Go back to your EAP vendor and ask for SUD-specific utilization, broken out by operating company. Not blended. Not all-conditions. Ask in writing.

Two things will happen. Either they can't produce it at that granularity, which tells you something — or they can, and the number will be well under one percent in most of your businesses. Either way, the gap between what gets reported to you and what's actually in your workforce is the whole business case, and you don't need me to make it.

Three — start one internal conversation about recovery-friendly practices. Leadership, HR, or Safety. It starts with a decision, not a budget.""")

# ---------------------------------------- slide 18 (NEW: pilot proposal)
s = S[17]
set_text(find(s, "Text 0"), "WHAT WE'RE PROPOSING")
set_text(find(s, "Text 1"), "Start with one operating company. 90 days.")

set_text(find(s, "Text 4"), "Pick the one company where the case is clearest.")
set_text(find(s, "Text 5"),
         "Highest-risk workforce, or simply the HR leader with the most "
         "appetite. One masthead, one business, one market.")
box(find(s, "Text 5"), width=10.60, height=0.30)

set_text(find(s, "Text 8"), "We run the full 90-day activation. Done for you.")
set_text(find(s, "Text 9"),
         "Communications, people-leader training, peer champions, the "
         "dedicated line and company code — built for that culture.")
box(find(s, "Text 9"), width=10.60, height=0.30)

set_text(find(s, "Text 12"),
         "You judge it on awareness and utilization at day 90.")
set_text(find(s, "Text 13"),
         "If it tracks Gardner and Olympia, we scale to the next company. If "
         "it doesn't, you have spent a quarter and learned something.")
box(find(s, "Text 13"), width=10.60, height=0.30)

set_text(find(s, "Text 14"),
         "The fastest way to know if this works in your workforce is to run "
         "it in one of them.")
drop(find(s, "Picture 16"))   # QR to the cost calculator: wrong slide for it

notes(s, """Jess

Here's what I'm actually asking for, and it is deliberately small.

Not a group-wide rollout. One operating company, ninety days.

Pick the one where the case is clearest to you — the highest-risk workforce, or honestly just whichever HR leader in this room has the most appetite. One masthead, one business, one market.

We run the whole activation. Communications, people-leader training, peer champions, the dedicated line with a code specific to that company. Your team's job is to give us your leaders for one session.

Then you judge it on two numbers at day ninety: awareness and utilization. If it tracks what Gardner and Olympia saw, we have a conversation about the next company. If it doesn't, you have spent a quarter and learned something real about your own workforce.

I'd rather earn the second company than be handed all of them.""")

# ------------------------------------------------ slide 19 (contact)
notes(S[18], """Jess

Questions. And the two I expect, so let me pre-empt them.

Global coverage — SOBRPath's resource directory is built out for all fifty US states. If we're piloting inside News UK or News Corp Australia, the peer coaching and the 24/7 navigation travel, but the local resource directory is work we'd scope with you before committing to a date. I'd rather say that now than discover it in month two. [Jess: confirm current non-US coverage with Marin before Friday.]

Privacy — anonymous by design. No names shared with HR, no claims filed, nothing hits the plan. That's also why utilization reporting comes back aggregated at the company level.""")

prs.save(OUT)
print("saved", OUT)
